#!/usr/bin/env python3
"""大きな文書での所要時間・最大メモリを測る（開発用。CI では実行しない）.

    pip install reportlab
    python scripts/bench.py                 # 既定サイズ
    python scripts/bench.py --scale 0.2     # 小さめで素早く

生成する文書（一時ディレクトリ。終了時に削除）:
- big.xlsx : 行数 N（既定 100,000 行 × 10 列）
- big.pdf  : ページ数 P（既定 300 ページ、1 ページ 40 行）
- big.docx : 段落 D（既定 5,000 段落 + 表 100 個）

各ファイルを別プロセスで doc2md 変換し、所要時間と最大常駐メモリ（ru_maxrss）を表で出す。
xlsx は「通常読み込み（--max-cells 無制限）」と「省メモリ読み込み（--max-cells 0）」の両方を測る。
"""

from __future__ import annotations

import argparse
import json
import subprocess
import sys
import tempfile
import time
from pathlib import Path

CHILD = r"""
import json, resource, sys, time
from pathlib import Path
from doc_to_markdown.converter import convert_file
from doc_to_markdown.model import Options
src, out, max_cells = Path(sys.argv[1]), Path(sys.argv[2]), int(float(sys.argv[3]))
t = time.perf_counter()
r = convert_file(src, out, options=Options(xlsx_full_load_max_cells=max_cells))
elapsed = time.perf_counter() - t
rss = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
rss_mb = rss / 1024 / 1024 if sys.platform == "darwin" else rss / 1024
print(json.dumps({"seconds": elapsed, "peak_mb": rss_mb, "md_mb": r.output.stat().st_size / 1e6,
                  "warnings": [f.rule_id for f in r.findings if f.severity.value == "warning"]}))
"""


def make_xlsx(path: Path, rows: int) -> None:
    import openpyxl

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("data")
    ws.append([f"列{i}" for i in range(1, 11)])
    for r in range(rows):
        ws.append(
            [
                f"ID-{r:07d}",
                r,
                r * 1.5,
                "区分A" if r % 2 else "区分B",
                "2026-04-01",
                r % 97,
                "備考テキスト",
                r * 3,
                "東京都",
                "OK",
            ]
        )
    wb.save(path)


def make_pdf(path: Path, pages: int) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    c = canvas.Canvas(str(path), pagesize=A4)
    _, h = A4
    for p in range(pages):
        c.setFont("HeiseiKakuGo-W5", 16)
        c.drawString(50, h - 60, f"第{p + 1}章 性能確認用の見出し")
        c.setFont("HeiseiKakuGo-W5", 10)
        for i in range(40):
            c.drawString(50, h - 90 - i * 17, f"これは性能確認のための本文です。行番号 {i + 1}。ページ {p + 1}。")
        c.showPage()
    c.save()


def make_docx(path: Path, paragraphs: int) -> None:
    import docx

    d = docx.Document()
    for i in range(paragraphs):
        if i % 50 == 0:
            d.add_heading(f"節 {i // 50 + 1}", level=1)
            t = d.add_table(rows=6, cols=4)
            for r in range(6):
                for c in range(4):
                    t.cell(r, c).text = f"r{r}c{c}"
        d.add_paragraph(f"これは性能確認のための段落 {i + 1} です。" * 3)
    d.save(path)


def run(src: Path, out: Path, max_cells: float) -> dict:
    r = subprocess.run(
        [sys.executable, "-c", CHILD, str(src), str(out), str(max_cells)], capture_output=True, text=True, check=True
    )
    return json.loads(r.stdout.strip().splitlines()[-1])


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--scale", type=float, default=1.0)
    args = ap.parse_args()
    rows, pages, paras = int(100_000 * args.scale), int(300 * args.scale), int(5_000 * args.scale)
    results = []
    with tempfile.TemporaryDirectory() as d:
        tmp = Path(d)
        t = time.perf_counter()
        make_xlsx(tmp / "big.xlsx", rows)
        make_pdf(tmp / "big.pdf", pages)
        make_docx(tmp / "big.docx", paras)
        print(f"generated in {time.perf_counter() - t:.1f}s", file=sys.stderr)
        cases = [
            ("big.xlsx", f"{rows:,} 行 × 10 列（通常読み込み）", 1e12),
            ("big.xlsx", f"{rows:,} 行 × 10 列（省メモリ読み込み）", 0),
            ("big.pdf", f"{pages} ページ", 200_000),
            ("big.docx", f"{paras:,} 段落 + 表 {paras // 50} 個", 200_000),
        ]
        for name, label, max_cells in cases:
            src = tmp / name
            res = run(src, tmp / f"out-{len(results)}", max_cells)
            results.append((name, label, src.stat().st_size / 1e6, res))
    print("| 入力 | 内容 | 入力サイズ | 所要時間 | 最大メモリ | 出力 .md | 警告 |")
    print("|---|---|---|---|---|---|---|")
    for name, label, size, r in results:
        print(
            f"| {name} | {label} | {size:.1f} MB | {r['seconds']:.1f} 秒 | {r['peak_mb']:.0f} MB | "
            f"{r['md_mb']:.1f} MB | {', '.join(r['warnings']) or '-'} |"
        )
    return 0


if __name__ == "__main__":
    sys.exit(main())
