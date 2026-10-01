"""Excel (.xlsx / .xlsm) → 中間表現（openpyxl）.

実務の Excel によくある形に対応する。

| パターン | 扱い |
|---|---|
| 表の上のタイトル行・作成者行 | 段落として表の前に出す（見出し行と取り違えない） |
| 見出し行が 1 行目でない | 自動検出（文字の入ったセルが多い最初の行）。`--header-row` で明示も可 |
| 2 段の見出し（「金額」の下に「税抜 / 税額 / 税込」） | 列名を「金額/税抜」のように連結 |
| セル結合 | 見出しは結合範囲の全列に同じ名前、本文の縦結合は値を繰り返す（検証しやすくするため） |
| 1 シートに複数の表 | 空行・空列で区切られた範囲ごとに別の表（`### 表N（A4:G11）`） |
| 表の下の注記（※…） | 表の後ろに段落として出す |
| 非表示の行・列・シート | 既定は出力して注記。`--skip-hidden` で除外 |
| 数式 | 保存されている計算結果を出す。結果が無いブックは警告（`--recalc` で LibreOffice 再計算） |
| パーセント表示のセル | `10%` のように表示形式どおりに出す |
| 大きなブック | 20 万セル（`--max-cells`）を超えると省メモリ読み込み（1 シート = 1 表、結合の解決なし） |
"""

from __future__ import annotations

import datetime as dt
import math
import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from devtools_common.tempfiles import secure_tempdir
from lxml import etree
from openpyxl.utils import get_column_letter

from ..assets import AssetWriter
from ..model import Block, Document, Heading, Options, Paragraph, StreamTable, Table, count_chars

NOTE_PREFIXES = ("※", "注", "備考", "＊", "*")
MAX_HEADER_DEPTH = 3


def format_value(v, number_format: str = "") -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)) and "%" in (number_format or ""):
        decimals = 0
        if "." in number_format:
            decimals = len(number_format.split(".", 1)[1].split("%", 1)[0].strip("_ )"))
        return f"{v * 100:.{decimals}f}%"
    if isinstance(v, float):
        if math.isnan(v) or math.isinf(v):
            return str(v)
        return str(int(v)) if v.is_integer() else repr(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    return str(v)


# ---------------------------------------------------------------------------
# シートの格子
# ---------------------------------------------------------------------------
@dataclass
class Merge:
    r0: int
    c0: int
    r1: int
    c1: int

    @property
    def width(self) -> int:
        return self.c1 - self.c0 + 1


class Grid:
    """1 シート分のセル値（文字列）と結合情報. 行・列は 1 始まり."""

    def __init__(self, ws, opts: Options):
        self.values: dict[tuple[int, int], str] = {}
        self.merge_at: dict[tuple[int, int], Merge] = {}  # 結合範囲内の全セル → 結合
        self.hidden_rows: set[int] = set()
        self.hidden_cols: set[int] = set()
        for r, dim in ws.row_dimensions.items():
            if dim.hidden:
                self.hidden_rows.add(r)
        for dim in ws.column_dimensions.values():
            if dim.hidden and dim.min and dim.max:
                self.hidden_cols.update(range(dim.min, dim.max + 1))
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                text = format_value(c.value, c.number_format)
                if text.strip():
                    self.values[(c.row, c.column)] = text
        for rng in ws.merged_cells.ranges:
            m = Merge(rng.min_row, rng.min_col, rng.max_row, rng.max_col)
            for r in range(m.r0, m.r1 + 1):
                for c in range(m.c0, m.c1 + 1):
                    self.merge_at[(r, c)] = m
        self.skip_hidden = opts.skip_hidden

    def visible(self, r: int, c: int) -> bool:
        return not self.skip_hidden or (r not in self.hidden_rows and c not in self.hidden_cols)

    def output_text(self) -> str:
        """出力対象のセル値（非表示を除外する設定ならそれを除く）. ロス検知の独立カウントに使う."""
        return "".join(v for (r, c), v in self.values.items() if self.visible(r, c))

    def occupied(self, r: int, c: int) -> bool:
        m = self.merge_at.get((r, c))
        if m:
            return (m.r0, m.c0) in self.values
        return (r, c) in self.values

    def value(self, r: int, c: int, *, fill_merged: bool) -> str:
        m = self.merge_at.get((r, c))
        if m and fill_merged:
            return self.values.get((m.r0, m.c0), "")
        if m and (r, c) != (m.r0, m.c0):
            return ""
        return self.values.get((r, c), "")

    def bounds(self) -> tuple[int, int, int, int] | None:
        cells = list(self.values) + [(m.r1, m.c1) for m in self.merge_at.values() if (m.r0, m.c0) in self.values]
        if not cells:
            return None
        return min(r for r, _ in cells), min(c for _, c in cells), max(r for r, _ in cells), max(c for _, c in cells)


def split_regions(g: Grid, r0: int, c0: int, r1: int, c1: int, depth: int = 0) -> list[tuple[int, int, int, int]]:
    """空行・空列で区切られた矩形に分割する（XY-cut）. 上→下、左→右の順."""

    def row_used(r: int) -> bool:
        return any(g.occupied(r, c) for c in range(c0, c1 + 1))

    def col_used(c: int, ra: int, rb: int) -> bool:
        return any(g.occupied(r, c) for r in range(ra, rb + 1))

    # 上下の空行を詰める
    while r0 <= r1 and not row_used(r0):
        r0 += 1
    while r1 >= r0 and not row_used(r1):
        r1 -= 1
    if r0 > r1:
        return []
    # 左右の空列を詰める
    while c0 <= c1 and not col_used(c0, r0, r1):
        c0 += 1
    while c1 >= c0 and not col_used(c1, r0, r1):
        c1 -= 1
    if depth > 6:
        return [(r0, c0, r1, c1)]
    # 空行で分割
    bands, start = [], r0
    for r in range(r0, r1 + 2):
        if r > r1 or not row_used(r):
            if start <= r - 1:
                bands.append((start, r - 1))
            start = r + 1
    if len(bands) > 1:
        return [x for a, b in bands for x in split_regions(g, a, c0, b, c1, depth + 1)]
    # 空列で分割
    cols, start = [], c0
    for c in range(c0, c1 + 2):
        if c > c1 or not col_used(c, r0, r1):
            if start <= c - 1:
                cols.append((start, c - 1))
            start = c + 1
    if len(cols) > 1:
        return [x for a, b in cols for x in split_regions(g, r0, a, r1, b, depth + 1)]
    return [(r0, c0, r1, c1)]


def _logical_cells(g: Grid, r: int, c0: int, c1: int) -> list[str]:
    """行 r の、結合を 1 つと数えた値の一覧（空は除く）."""
    out, seen = [], set()
    for c in range(c0, c1 + 1):
        m = g.merge_at.get((r, c))
        key = (m.r0, m.c0) if m else (r, c)
        if key in seen:
            continue
        seen.add(key)
        v = g.values.get(key, "") if m else g.values.get((r, c), "")
        if v and g.visible(*key):
            out.append(v)
    return out


def _is_number(s: str) -> bool:
    try:
        float(s.replace(",", "").rstrip("%"))
        return True
    except ValueError:
        return False


def detect_header_row(g: Grid, r0: int, c0: int, r1: int, c1: int) -> int | None:
    width = c1 - c0 + 1
    need = max(2, math.ceil(width * 0.5))
    for r in range(r0, min(r1, r0 + 20) + 1):
        vals = _logical_cells(g, r, c0, c1)
        filled = sum(1 for c in range(c0, c1 + 1) if g.occupied(r, c))
        if filled >= need and len(vals) >= 2 and sum(1 for v in vals if not _is_number(v)) >= len(vals) * 0.7:
            return r
    return None


def header_depth(g: Grid, h: int, c0: int, c1: int, r1: int) -> int:
    """見出しが何行あるか: h 行目の結合が下の行まで伸びている、または横結合の下に子見出しがある."""
    depth = 1

    def looks_like_subheader(r: int) -> bool:
        """r 行目の、上からの結合で覆われていないセルがすべて文字（数値でない）で、1 つ以上ある."""
        free = [g.values.get((r, c), "") for c in range(c0, c1 + 1) if not ((m := g.merge_at.get((r, c))) and m.r0 < r)]
        filled = [v for v in free if v]
        return bool(filled) and not any(_is_number(v) for v in filled)

    for c in range(c0, c1 + 1):
        m = g.merge_at.get((h, c))
        if not m or m.r0 != h:
            continue
        if m.r1 > h and looks_like_subheader(h + 1):  # 縦結合（A4:A5）の横に子見出しがある
            depth = max(depth, m.r1 - h + 1)
        elif m.width > 1 and h + 1 <= r1:  # 横結合の下に子見出しがあるか
            below = [g.values.get((h + 1, cc), "") for cc in range(m.c0, m.c1 + 1)]
            if all(below) and not any(_is_number(b) for b in below):
                depth = max(depth, 2)
    return min(depth, MAX_HEADER_DEPTH, r1 - h + 1)


def header_names(g: Grid, h: int, depth: int, c0: int, c1: int) -> list[str]:
    names = []
    for c in range(c0, c1 + 1):
        parts: list[str] = []
        for r in range(h, h + depth):
            v = g.value(r, c, fill_merged=True).replace("\n", " ")
            if v and (not parts or parts[-1] != v):
                parts.append(v)
        names.append("/".join(parts))
    return names


@dataclass
class SheetResult:
    blocks: list[Block]
    tables: int
    rows: int
    text: str
    filled: int


def sheet_blocks(ws, g: Grid, opts: Options, sheet_title: str) -> SheetResult:
    blocks: list[Block] = []
    b = g.bounds()
    if b is None:
        return SheetResult([Paragraph("（空のシート）", synthetic=True)], 0, 0, "", 0)
    regions = split_regions(g, *b)
    forced = opts.header_rows.get(sheet_title, opts.header_rows.get("*"))
    table_regions = []
    for reg in regions:
        r0, c0, r1, c1 = reg
        if forced and r0 <= forced <= r1:
            h = forced
        elif forced:
            h = None
        else:
            h = detect_header_row(g, r0, c0, r1, c1) if (c1 > c0) else None
        table_regions.append((reg, h))
    n_tables = sum(1 for _, h in table_regions if h is not None)

    count_tables = rows_total = filled = 0
    texts: list[str] = []
    for (r0, c0, r1, c1), h in table_regions:
        if h is None:  # 表でない範囲（タイトル・作成者・注記）
            for r in range(r0, r1 + 1):
                if opts.skip_hidden and r in g.hidden_rows:
                    continue
                vals = _logical_cells(g, r, c0, c1)
                if vals:
                    line = " / ".join(vals)
                    texts.append(line)
                    blocks.append(Paragraph(line))
            continue
        count_tables += 1
        # 見出し行より上（同じ範囲内）はタイトルとして段落に
        for r in range(r0, h):
            vals = _logical_cells(g, r, c0, c1)
            if vals:
                line = " / ".join(vals)
                texts.append(line)
                blocks.append(Paragraph(line))
        depth = header_depth(g, h, c0, c1, r1)
        header = header_names(g, h, depth, c0, c1)
        texts += [v for r in range(h, h + depth) for v in _logical_cells(g, r, c0, c1)]
        body_rows: list[list[str]] = []
        notes: list[str] = []
        last = r1
        # 表の末尾の注記行（1 セルだけで ※ 等で始まる）
        while last > h + depth - 1:
            vals = _logical_cells(g, last, c0, c1)
            if len(vals) == 1 and vals[0].lstrip().startswith(NOTE_PREFIXES):
                notes.insert(0, vals[0])
                last -= 1
            else:
                break
        for r in range(h + depth, last + 1):
            if opts.skip_hidden and r in g.hidden_rows:
                continue
            row = []
            for c in range(c0, c1 + 1):
                m = g.merge_at.get((r, c))
                vertical = m is not None and m.r1 > m.r0 and m.c0 == c and (r, c) != (m.r0, m.c0)
                v = g.value(r, c, fill_merged=vertical)
                if vertical and v:
                    filled += 1
                row.append(v)
            if any(row):
                texts += [v for v in _logical_cells(g, r, c0, c1)]
                body_rows.append(row)
        if opts.skip_hidden:
            keep = [i for i, c in enumerate(range(c0, c1 + 1)) if c not in g.hidden_cols]
            header = [header[i] for i in keep]
            body_rows = [[row[i] for i in keep] for row in body_rows]
        if n_tables > 1:
            ref = f"{get_column_letter(c0)}{h}:{get_column_letter(c1)}{last}"
            blocks.append(Heading(3, f"表{count_tables}（{ref}）", synthetic=True))
        blocks.append(Table(rows=[header, *body_rows], header=True))
        rows_total += 1 + len(body_rows)
        for n in notes:
            texts.append(n)
            blocks.append(Paragraph(n))
    return SheetResult(blocks, count_tables, rows_total, "".join(texts), filled)


# ---------------------------------------------------------------------------
# 省メモリ（大きなブック）
# ---------------------------------------------------------------------------
def _stream_sheet(ws, doc: Document) -> Iterator[Block]:
    it = ws.iter_rows(values_only=True)
    head: list[list[str]] = []
    for values in it:
        cells = [format_value(v) for v in values]
        if any(c.strip() for c in cells):
            head.append(cells)
            if len(head) >= 20:
                break
    if not head:
        yield Paragraph("（空のシート）", synthetic=True)
        return
    width = max(len(r) for r in head)
    header_idx = 0
    for i, r in enumerate(head):
        texts = [c for c in r if c.strip()]
        if (
            len(texts) >= max(2, math.ceil(width * 0.5))
            and sum(1 for c in texts if not _is_number(c)) >= len(texts) * 0.7
        ):
            header_idx = i
            break
    for r in head[:header_idx]:
        line = " / ".join(c for c in r if c.strip())
        doc.stats.add_text(line)
        yield Paragraph(line)
    st = doc.stats
    st.tables += 1

    def rows():
        for r in head[header_idx:]:
            st.table_rows += 1
            st.add_text("".join(r))
            yield r
        for values in it:
            cells = [format_value(v) for v in values]
            if any(c.strip() for c in cells):
                st.table_rows += 1
                st.add_text("".join(cells))
                yield cells

    gen = rows()
    header = next(gen)
    yield StreamTable(header=header, rows=gen)


# ---------------------------------------------------------------------------
_SHEET = re.compile(r"xl/worksheets/sheet\d+\.xml$")
NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def scan_workbook(path: Path) -> tuple[int, int]:
    """(セル数, 計算結果が保存されていない数式セルの数) を、シートの XML を 1 回流し読みして数える.

    openpyxl で読むより速く、メモリも一定（大きなブックの判定に使う）。
    """
    cells = uncached = 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not _SHEET.match(name):
                continue
            with z.open(name) as f:
                for _, row in etree.iterparse(
                    f, events=("end",), tag=f"{NS_MAIN}row", resolve_entities=False, no_network=True, huge_tree=True
                ):
                    for c in row.iterchildren(f"{NS_MAIN}c"):
                        cells += 1
                        if c.find(f"{NS_MAIN}f") is not None:
                            v = c.find(f"{NS_MAIN}v")
                            if v is None or v.text is None:
                                uncached += 1
                    # 読み終えた行を木から外してメモリを一定に保つ
                    row.clear()
                    parent = row.getparent()
                    while row.getprevious() is not None:
                        del parent[0]
    return cells, uncached


def _blocks(path: Path, doc: Document, opts: Options, cells: int) -> Iterator[Block]:
    st = doc.stats
    yield Heading(1, path.stem, synthetic=True)
    full = cells <= opts.xlsx_full_load_max_cells
    wb = openpyxl.load_workbook(str(path), read_only=not full, data_only=True)
    if not full:
        doc.notes.append(
            f"大きなブック（{cells:,} セル）のため省メモリで読み込み: "
            "セル結合・複数の表・非表示の判定は行わず、1 シート = 1 表で出力"
        )
    raw = 0
    hidden_rows = hidden_sheets = filled = 0
    try:
        for ws in wb.worksheets:
            hidden = getattr(ws, "sheet_state", "visible") != "visible"
            if hidden:
                hidden_sheets += 1
                if opts.skip_hidden:
                    continue
            st.headings += 1
            yield Heading(2, ws.title + ("（非表示）" if hidden else ""))
            if not full:
                yield from _stream_sheet(ws, doc)
                continue
            g = Grid(ws, opts)
            raw += count_chars(g.output_text())
            hidden_rows += len([r for r in g.hidden_rows if any(k[0] == r for k in g.values)])
            res = sheet_blocks(ws, g, opts, ws.title)
            st.tables += res.tables
            st.table_rows += res.rows
            st.add_text(res.text)
            filled += res.filled
            yield from res.blocks
    finally:
        wb.close()
    if full:
        doc.raw_text_chars = raw
    if hidden_sheets:
        doc.notes.append(f"非表示のシート {hidden_sheets} 枚" + ("を除外" if opts.skip_hidden else "を含めて出力"))
    if hidden_rows and not opts.skip_hidden:
        doc.warnings.append(f"非表示の行 {hidden_rows} 行を含めて出力しています（除外するには --skip-hidden）")
    if filled:
        doc.notes.append(f"縦に結合されたセル {filled} 個は上の値を繰り返して出力")


def convert(path: Path, assets: AssetWriter, opts: Options | None = None) -> Document:
    opts = opts or Options()
    doc = Document()
    cells, uncached = scan_workbook(path)
    if uncached and opts.recalc:
        doc.notes.append(f"数式の計算結果が無いセル {uncached} 個を LibreOffice で再計算して出力")
        doc.blocks = _recalc_blocks(path, doc, opts, cells)
        return doc
    if uncached:
        doc.warnings.append(
            f"数式の計算結果が保存されていないセルが {uncached} 個あり、空欄で出力しています"
            "（Excel で開いて保存し直すか、--recalc で LibreOffice に再計算させてください）"
        )
    doc.blocks = _blocks(path, doc, opts, cells)
    return doc


def _recalc_blocks(path: Path, doc: Document, opts: Options, cells: int) -> Iterator[Block]:
    from ..office import recalc_xlsx

    with secure_tempdir("doc2md-recalc-") as tmp:
        fixed = recalc_xlsx(path, tmp)
        for b in _blocks(fixed, doc, opts, cells):
            if isinstance(b, Heading) and b.synthetic and b.level == 1:
                b = Heading(1, path.stem, synthetic=True)
            yield b
