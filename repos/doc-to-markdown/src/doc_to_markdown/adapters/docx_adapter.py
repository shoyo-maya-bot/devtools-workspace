"""Word (.docx) → 中間表現.

python-docx の高水準 API は段落直下の文字しか見ないため（ハイパーリンク内・変更履歴内・コンテンツ
コントロール内・テキストボックス内の文字を落とす）、本文は XML を直接たどって組み立てる。

対応している要素
- 見出し（Title / Heading N / アウトラインレベル）、段落、太字・斜体、ハイパーリンク `[text](url)`
- 箇条書き（番号付き / 記号付き、レベル）
- 表: 横結合（gridSpan）、縦結合（vMerge: 値を繰り返して出力）、入れ子の表（セル内に行を展開）、セル内の画像
- テキストボックス・図形内の文字（`> **テキストボックス**: …`）
- 画像（本文・表・テキストボックス内）
- 脚注・文末脚注（GFM の `[^1]` / `[^e1]`）
- コメント（`[^c1]: 作成者: 内容（対象:「…」）`）
- 変更履歴（--track-changes accept / reject / show）
- コンテンツコントロール（w:sdt）・フィールド・スマートタグの中の文字
- ヘッダー・フッター（重複を除いて末尾にまとめる）
- 数式（テキストとして出力）
"""

from __future__ import annotations

import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as DocxParagraph

from ..assets import AssetWriter
from ..model import (
    Block,
    Document,
    FootnoteDef,
    Heading,
    Image,
    ListItem,
    Options,
    Paragraph,
    Quote,
    SourceStats,
    Table,
    count_chars,
)

MC = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
VML = "{urn:schemas-microsoft-com:vml}"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
M_T = "{http://schemas.openxmlformats.org/officeDocument/2006/math}t"

_HEADING_STYLE = re.compile(r"^(?:heading|見出し)\s*(\d)$", re.I)
_LIST_STYLE = re.compile(r"^List (Bullet|Number)(?: (\d))?$")  # 組み込みの箇条書きスタイル（"List Bullet 2" = 2 段目）
_FALSE = {"0", "false", "off", "none"}

T = qn("w:t")
DEL_T = qn("w:delText")
R_TAG = qn("w:r")
P_TAG = qn("w:p")
TBL = qn("w:tbl")
TC = qn("w:tc")
TR = qn("w:tr")
SDT = qn("w:sdt")
SDT_CONTENT = qn("w:sdtContent")
TXBX = qn("w:txbxContent")
INS_TAGS = {qn("w:ins"), qn("w:moveTo")}
DEL_TAGS = {qn("w:del"), qn("w:moveFrom")}
TRANSPARENT = {qn("w:smartTag"), qn("w:customXml"), qn("w:fldSimple"), qn("w:dir"), qn("w:bdo")}


def _on(el) -> bool:
    return el is not None and (el.get(qn("w:val")) or "true").lower() not in _FALSE


def _has_ancestor(el, tags: set[str], stop=None) -> bool:
    p = el.getparent()
    while p is not None and p is not stop:
        if p.tag in tags:
            return True
        p = p.getparent()
    return False


# ---------------------------------------------------------------------------
# 独立経路の文字数カウント（ロス検知用）
# ---------------------------------------------------------------------------
def raw_text_chars(path: Path, opts: Options) -> int:
    """zip 内の XML から直接 w:t を数える（アダプタの解釈を通さない）."""
    fallback = {f"{MC}Fallback"}
    total = 0
    seen_hf: set[str] = set()
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if n.startswith("word/") and n.endswith(".xml")]
        for name in sorted(names):
            base = name.rsplit("/", 1)[-1]
            is_hf = re.match(r"(header|footer)\d*\.xml$", base)
            if not (
                base in ("document.xml", "footnotes.xml", "endnotes.xml")
                or is_hf
                or (base == "comments.xml" and opts.include_comments)
            ):
                continue
            if is_hf and not opts.include_header_footer:
                continue
            root = parse_xml(z.read(name))
            parts: list[str] = []
            for el in root.iter(T, DEL_T):
                if _has_ancestor(el, fallback):
                    continue
                in_del = el.tag == DEL_T or _has_ancestor(el, DEL_TAGS)
                in_ins = _has_ancestor(el, INS_TAGS)
                if opts.track_changes == "accept" and in_del:
                    continue
                if opts.track_changes == "reject" and in_ins:
                    continue
                parts.append(el.text or "")
            text = "".join(parts)
            if is_hf:
                key = re.sub(r"\s+", "", text)
                if key in seen_hf:
                    continue
                seen_hf.add(key)
            total += count_chars(text)
    return total


# ---------------------------------------------------------------------------
# 段落内（インライン）の描画
# ---------------------------------------------------------------------------
@dataclass
class _Inline:
    text: str = ""
    images: list[str] = field(default_factory=list)
    textboxes: list[str] = field(default_factory=list)


class _Converter:
    def __init__(self, path: Path, assets: AssetWriter, opts: Options):
        self.path, self.assets, self.opts = path, assets, opts
        self.document = docx.Document(str(path))
        self.part = self.document.part
        self.doc = Document()
        self.st: SourceStats = self.doc.stats
        self.formats = self._numbering_formats()
        self.style_names, self.default_style = self._style_map()
        self.title_offset = 1 if any(self._is_title(p) for p in self.document.paragraphs) else 0
        self.notes = {
            "footnote": self._load_notes(RT.FOOTNOTES, "w:footnote"),
            "endnote": self._load_notes(RT.ENDNOTES, "w:endnote"),
        }
        self.comments = self._load_notes(RT.COMMENTS, "w:comment") if opts.include_comments else {}
        self.labels: dict[tuple[str, str], str] = {}  # (種類, id) → 表示ラベル（出現順）
        self.comment_capture: dict[str, list[str]] = {}
        self.active_comments: set[str] = set()
        self.revisions = 0
        self.vmerge_filled = 0
        self.nested_tables = 0
        self.linked_images = 0

    # -- パーツ -------------------------------------------------------------
    def _related_element(self, reltype: str):
        for rel in self.part.rels.values():
            if rel.reltype == reltype and not rel.is_external:
                tp = rel.target_part
                return getattr(tp, "element", None) if hasattr(tp, "element") else parse_xml(tp.blob)
        return None

    def _load_notes(self, reltype: str, tag: str) -> dict[str, object]:
        root = self._related_element(reltype)
        if root is None:
            return {}
        out = {}
        for el in root.iter(qn(tag)):
            if el.get(qn("w:type")) in ("separator", "continuationSeparator", "continuationNotice"):
                continue
            out[el.get(qn("w:id"))] = el
        return out

    def _numbering_formats(self) -> dict[tuple[str, str], str]:
        try:
            numbering = self.part.numbering_part.element
        except (KeyError, NotImplementedError, AttributeError):
            return {}
        abstract: dict[str, dict[str, str]] = {}
        for an in numbering.findall(qn("w:abstractNum")):
            levels = {}
            for lvl in an.findall(qn("w:lvl")):
                fmt = lvl.find(qn("w:numFmt"))
                levels[lvl.get(qn("w:ilvl"))] = fmt.get(qn("w:val")) if fmt is not None else "bullet"
            abstract[an.get(qn("w:abstractNumId"))] = levels
        result = {}
        for num in numbering.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            if ref is None:
                continue
            for ilvl, fmt in abstract.get(ref.get(qn("w:val")), {}).items():
                result[(num.get(qn("w:numId")), ilvl)] = fmt
        return result

    # -- 段落の種類 -----------------------------------------------------------
    def _style_map(self) -> tuple[dict[str, str], str]:
        """スタイル ID → 表示名（python-docx の p.style は段落ごとに全スタイルを走査して遅いため一度だけ作る）."""
        names: dict[str, str] = {}
        default = ""
        for s in self.document.styles:
            try:
                names[s.style_id] = s.name or ""
                if s.type == 1 and s.element.get(qn("w:default")) in ("1", "true"):  # 1 = 段落スタイル
                    default = s.name or ""
            except (KeyError, ValueError, AttributeError):
                continue
        return names, default

    def _style_name(self, p: DocxParagraph) -> str:
        ppr = p._p.pPr
        ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
        if ps is None:
            return self.default_style
        return self.style_names.get(ps.get(qn("w:val")), "")

    def _is_title(self, p: DocxParagraph) -> bool:
        return self._style_name(p).lower() == "title"

    def _heading_level(self, p: DocxParagraph) -> int | None:
        if self._is_title(p):
            return 1
        m = _HEADING_STYLE.match(self._style_name(p).strip())
        if m:
            return int(m.group(1)) + self.title_offset
        ppr = p._p.pPr
        if ppr is not None:
            ol = ppr.find(qn("w:outlineLvl"))
            if ol is not None:
                val = int(ol.get(qn("w:val")))
                if val < 9:
                    return val + 1 + self.title_offset
        return None

    def _list_info(self, p: DocxParagraph) -> tuple[int, bool] | None:
        ppr = p._p.pPr
        numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
        name = self._style_name(p)
        style_list = _LIST_STYLE.match(name)
        style_level = int(style_list.group(2) or 1) - 1 if style_list else 0
        if numpr is None:
            if style_list:
                return style_level, style_list.group(1) == "Number"
            return None
        ilvl_el, numid_el = numpr.find(qn("w:ilvl")), numpr.find(qn("w:numId"))
        ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else str(style_level)
        num_id = numid_el.get(qn("w:val")) if numid_el is not None else None
        if num_id == "0":  # 番号付け解除
            return None
        fmt = self.formats.get((num_id, ilvl), "bullet" if "Bullet" in name else "decimal")
        return int(ilvl), fmt not in ("bullet", "none")

    # -- インライン -----------------------------------------------------------
    def _label(self, kind: str, note_id: str) -> str:
        key = (kind, note_id)
        if key not in self.labels:
            prefix = {"footnote": "", "endnote": "e", "comment": "c"}[kind]
            n = sum(1 for k in self.labels if k[0] == kind) + 1
            self.labels[key] = f"{prefix}{n}"
        return self.labels[key]

    def _image(self, blip, out: _Inline) -> None:
        rid = blip.get(qn("r:embed"))
        if not rid:
            if blip.get(qn("r:link")):
                self.linked_images += 1
            return
        part = self.part.related_parts.get(rid)
        if part is not None and hasattr(part, "blob"):
            out.images.append(self.assets.write(part.blob, content_type=getattr(part, "content_type", "")))

    def _textbox(self, txbx, out: _Inline) -> None:
        texts = [self.inline(p).text for p in txbx.iter(P_TAG) if not _has_ancestor(p, {TXBX}, stop=txbx)]
        text = "\n".join(t for t in texts if t.strip())
        if text:
            out.textboxes.append(text)

    def _drawing(self, el, out: _Inline) -> None:
        """w:drawing / w:pict: 画像とテキストボックス. 入れ子のテキストボックスは外側だけ処理する."""
        for sub in el.iter():
            if sub.tag == TXBX and not _has_ancestor(sub, {TXBX}, stop=el):
                self._textbox(sub, out)
            elif sub.tag == qn("a:blip") and not _has_ancestor(sub, {TXBX}, stop=el):
                self._image(sub, out)
            elif sub.tag == f"{VML}imagedata" and not _has_ancestor(sub, {TXBX}, stop=el):
                rid = sub.get(qn("r:id"))
                part = self.part.related_parts.get(rid) if rid else None
                if part is not None and hasattr(part, "blob"):
                    out.images.append(self.assets.write(part.blob, content_type=getattr(part, "content_type", "")))

    def inline(self, p_el) -> _Inline:
        out = _Inline()
        buf: list[str] = []  # 同じ書式の連続した文字
        fmt: tuple[bool, bool] | None = None
        pieces: list[str] = []

        def emit_text(text: str, f: tuple[bool, bool]) -> None:
            nonlocal fmt
            if f != fmt:
                flush()
                fmt = f
            buf.append(text)
            for cid in self.active_comments:
                self.comment_capture.setdefault(cid, []).append(text)

        def flush() -> None:
            nonlocal fmt
            if buf:
                s = "".join(buf)
                b, i = fmt or (False, False)
                core = s.strip()
                if core and (b or i):
                    mark = "***" if b and i else "**" if b else "*"
                    lead, trail = s[: len(s) - len(s.lstrip())], s[len(s.rstrip()) :]
                    pieces.append(f"{lead}{mark}{core}{mark}{trail}")
                else:
                    pieces.append(s)
                buf.clear()
            fmt = None

        def emit_raw(token: str) -> None:
            flush()
            pieces.append(token)

        def walk(el, mode: str = "") -> None:  # mode: "" / "ins" / "del"
            for c in el:
                tag = c.tag
                if tag == R_TAG:
                    run(c, mode)
                elif tag == qn("w:hyperlink"):
                    sub = self.inline_fragment(c, mode)
                    rid = c.get(qn("r:id"))
                    url = ""
                    if rid and rid in self.part.rels and self.part.rels[rid].is_external:
                        url = self.part.rels[rid].target_ref
                    if sub.strip():  # コメント対象の文字は内側の描画で記録済み
                        emit_raw(f"[{sub.strip()}]({url})" if url else sub)
                elif tag in INS_TAGS:
                    self.revisions += 1
                    if self.opts.track_changes == "reject":
                        continue
                    if self.opts.track_changes == "show":
                        emit_raw("<ins>")
                        walk(c, "ins")
                        emit_raw("</ins>")
                    else:
                        walk(c, "ins")
                elif tag in DEL_TAGS:
                    self.revisions += 1
                    if self.opts.track_changes == "accept":
                        continue
                    if self.opts.track_changes == "show":
                        emit_raw("~~")
                        walk(c, "del")
                        emit_raw("~~")
                    else:
                        walk(c, "del")
                elif tag == SDT:
                    content = c.find(SDT_CONTENT)
                    if content is not None:
                        walk(content, mode)
                elif tag in TRANSPARENT:
                    walk(c, mode)
                elif tag == qn("w:commentRangeStart"):
                    if self.opts.include_comments:
                        self.active_comments.add(c.get(qn("w:id")))
                elif tag == qn("w:commentRangeEnd"):
                    self.active_comments.discard(c.get(qn("w:id")))
                elif tag == f"{MC}AlternateContent":
                    choice = c.find(f"{MC}Choice")
                    walk(choice if choice is not None else c, mode)
                elif tag.startswith("{http://schemas.openxmlformats.org/officeDocument/2006/math}"):
                    math = "".join(t.text or "" for t in c.iter(M_T))
                    if math:
                        emit_raw(math)

        def run(r, mode: str) -> None:
            rpr = r.find(qn("w:rPr"))
            f = (
                _on(rpr.find(qn("w:b"))) if rpr is not None else False,
                _on(rpr.find(qn("w:i"))) if rpr is not None else False,
            )
            for c in r:
                tag = c.tag
                if tag == T or tag == DEL_T and mode == "del":
                    emit_text(c.text or "", f)
                elif tag == qn("w:tab"):
                    emit_text(" ", f)
                elif tag in (qn("w:br"), qn("w:cr")):
                    if c.get(qn("w:type")) != "page":
                        emit_text("\n", f)
                elif tag == qn("w:noBreakHyphen"):
                    emit_text("-", f)
                elif tag in (qn("w:drawing"), qn("w:pict"), qn("w:object")):
                    flush()
                    self._drawing(c, out)
                elif tag == f"{MC}AlternateContent":
                    flush()
                    choice = c.find(f"{MC}Choice")
                    self._drawing(choice if choice is not None else c, out)
                elif tag == qn("w:footnoteReference"):
                    emit_raw(f"[^{self._label('footnote', c.get(qn('w:id')))}]")
                elif tag == qn("w:endnoteReference"):
                    emit_raw(f"[^{self._label('endnote', c.get(qn('w:id')))}]")
                elif tag == qn("w:commentReference"):
                    cid = c.get(qn("w:id"))
                    if self.opts.include_comments and cid in self.comments:
                        emit_raw(f"[^{self._label('comment', cid)}]")

        walk(p_el)
        flush()
        # 隣り合う変更履歴の記号をまとめる（~~8000 ~~~~円~~ → ~~8000 円~~）
        out.text = "".join(pieces).replace("~~~~", "").replace("</ins><ins>", "")
        return out

    def inline_fragment(self, el, mode: str) -> str:
        """ハイパーリンク等の内側だけを描画した文字列（画像・テキストボックスは本体へ回す）."""
        wrapper = parse_xml(f'<w:p xmlns:w="{W_NS}"/>')
        for c in el:
            wrapper.append(_clone(c))
        return self.inline(wrapper).text

    # -- ブロック -------------------------------------------------------------
    def blocks(self) -> Iterator[Block]:
        body = self.document.element.body
        for child in body.iterchildren():
            yield from self.block(child)
        if self.opts.include_header_footer:
            yield from self.header_footer()
        yield from self.note_definitions()
        self._finish_notes()

    def block(self, el) -> Iterator[Block]:
        tag = el.tag
        if tag == P_TAG:
            yield from self.paragraph(el)
        elif tag == TBL:
            yield self.table(el)
        elif tag == SDT:
            content = el.find(SDT_CONTENT)
            if content is not None:
                for c in content.iterchildren():
                    yield from self.block(c)
        elif tag in (qn("w:customXml"),):
            for c in el.iterchildren():
                yield from self.block(c)

    def paragraph(self, el) -> Iterator[Block]:
        p = DocxParagraph(el, self.document)
        inl = self.inline(el)
        text = inl.text
        if text.strip():
            self.st.add_text(_strip_markup(text))
            level = self._heading_level(p)
            if level:
                self.st.headings += 1
                yield Heading(level, text)
            else:
                info = self._list_info(p)
                if info:
                    self.st.list_items += 1
                    yield ListItem(text, info[0], info[1])
                else:
                    yield Paragraph(text)
        for tb in inl.textboxes:
            self.st.add_text(_strip_markup(tb))
            yield Quote(tb, "テキストボックス")
        for img in inl.images:
            self.st.images += 1
            yield Image(img, alt=Path(img).stem)

    def cell_text(self, tc) -> str:
        parts: list[str] = []
        for c in tc.iterchildren():
            if c.tag == P_TAG:
                inl = self.inline(c)
                if inl.text.strip():
                    parts.append(inl.text)
                parts += inl.textboxes
                parts += [f"![{Path(i).stem}]({i})" for i in inl.images]
            elif c.tag == TBL:
                self.nested_tables += 1
                rows = self.table_rows(c)
                parts += [" / ".join(x for x in r if x) for r in rows if any(r)]
            elif c.tag == SDT:
                content = c.find(SDT_CONTENT)
                if content is not None:
                    parts.append(self.cell_text(content))
        return "\n".join(parts)

    def table_rows(self, tbl) -> list[list[str]]:
        rows: list[list[str]] = []
        above: dict[int, str] = {}  # 列位置 → 直前行の値（縦結合の継続用）
        for tr in _iter_rows(tbl):
            row: list[str] = []
            trpr = tr.find(qn("w:trPr"))
            before = trpr.find(qn("w:gridBefore")) if trpr is not None else None
            if before is not None:
                row += [""] * int(before.get(qn("w:val")) or 0)
            for tc in _iter_cells(tr):
                tcpr = tc.find(qn("w:tcPr"))
                span = 1
                vmerge = None
                if tcpr is not None:
                    gs = tcpr.find(qn("w:gridSpan"))
                    span = int(gs.get(qn("w:val"))) if gs is not None else 1
                    vm = tcpr.find(qn("w:vMerge"))
                    if vm is not None:
                        vmerge = vm.get(qn("w:val")) or "continue"
                col = len(row)
                if vmerge == "continue":
                    value = above.get(col, "")
                    if value:
                        self.vmerge_filled += 1
                else:
                    value = self.cell_text(tc)
                row.append(value)
                row += [""] * (span - 1)
                above[col] = value
            rows.append(row)
        return rows

    def table(self, tbl) -> Table:
        rows = self.table_rows(tbl)
        self.st.tables += 1
        self.st.table_rows += sum(1 for r in rows if any(c.strip() for c in r))
        for r in rows:
            for c in r:
                self.st.add_text(_strip_markup(c))
        return Table(rows=rows, header=True)

    def header_footer(self) -> Iterator[Block]:
        items: list[str] = []
        seen: set[str] = set()
        order = {RT.HEADER: 0, RT.FOOTER: 1}
        rels = sorted(
            (r for r in self.part.rels.values() if not r.is_external and r.reltype in order),
            key=lambda r: (order[r.reltype], str(r.target_ref)),
        )
        for rel in rels:
            if rel.is_external or rel.reltype not in (RT.HEADER, RT.FOOTER):
                continue
            root = getattr(rel.target_part, "element", None)
            if root is None:
                continue
            label = "ヘッダー" if rel.reltype == RT.HEADER else "フッター"
            lines: list[str] = []
            for c in root.iterchildren():
                if c.tag == P_TAG:
                    inl = self.inline(c)
                    lines += [t for t in [inl.text, *inl.textboxes] if t.strip()]
                elif c.tag == TBL:
                    lines += [" / ".join(x for x in r if x) for r in self.table_rows(c) if any(r)]
            text = " / ".join(x.strip() for x in lines)
            key = re.sub(r"\s+", "", text)
            if text and key not in seen:
                seen.add(key)
                items.append(f"{label}: {text}")
        if items:
            yield Heading(2, "ヘッダー・フッター", synthetic=True)
            for it in items:
                self.st.list_items += 1
                self.st.add_text(it.split(": ", 1)[1])
                yield ListItem(it)

    def note_definitions(self) -> Iterator[Block]:
        for (kind, note_id), label in list(self.labels.items()):
            if kind in ("footnote", "endnote"):
                el = self.notes[kind].get(note_id)
                if el is None:
                    continue
                text = "\n".join(self.inline(p).text for p in el.iter(P_TAG))
                self.st.add_text(_strip_markup(text))
                yield FootnoteDef(label, text.strip())
            elif kind == "comment":
                el = self.comments.get(note_id)
                if el is None:
                    continue
                author = el.get(qn("w:author")) or "不明"
                body = " ".join(self.inline(p).text.strip() for p in el.iter(P_TAG))
                self.st.add_text(_strip_markup(body))
                target = "".join(self.comment_capture.get(note_id, [])).strip()
                suffix = f"（対象:「{target}」）" if target else ""
                yield FootnoteDef(label, f"コメント（{author}）: {body}{suffix}")

    def _finish_notes(self) -> None:
        if self.revisions:
            mode = {"accept": "承認した状態", "reject": "元に戻した状態", "show": "差分を記号付き"}[
                self.opts.track_changes
            ]
            self.doc.warnings.append(
                f"未確定の変更履歴が {self.revisions} 件あります（{mode}で出力。--track-changes で切替）"
            )
        n_comments = sum(1 for k in self.labels if k[0] == "comment")
        if n_comments:
            self.doc.notes.append(f"コメント {n_comments} 件を脚注形式（[^c1] …）で出力")
        if self.vmerge_filled:
            self.doc.notes.append(f"表の縦結合セル {self.vmerge_filled} 個は、上のセルの値を繰り返して出力")
        if self.nested_tables:
            self.doc.notes.append(f"入れ子の表 {self.nested_tables} 個はセル内に「列 / 列」形式で展開")
        if self.linked_images:
            self.doc.warnings.append(f"外部ファイルにリンクされた画像 {self.linked_images} 個は取り込めません")


def _clone(el):
    from copy import deepcopy

    return deepcopy(el)


def _iter_rows(tbl):
    for c in tbl.iterchildren():
        if c.tag == TR:
            yield c
        elif c.tag == SDT:
            content = c.find(SDT_CONTENT)
            if content is not None:
                yield from _iter_rows(content)


def _iter_cells(tr):
    for c in tr.iterchildren():
        if c.tag == TC:
            yield c
        elif c.tag == SDT:
            content = c.find(SDT_CONTENT)
            if content is not None:
                yield from _iter_cells(content)


_MD_MARKUP = re.compile(r"\[\^[^\]]+\]|\*\*\*|\*\*|\*|~~|</?ins>|!\[[^\]]*\]\([^)]+\)")
_MD_LINK = re.compile(r"\[([^\]]*)\]\([^)]+\)")


def _strip_markup(text: str) -> str:
    return _MD_MARKUP.sub("", _MD_LINK.sub(r"\1", text))


def convert(path: Path, assets: AssetWriter, opts: Options | None = None) -> Document:
    opts = opts or Options()
    conv = _Converter(path, assets, opts)
    doc = conv.doc
    doc.blocks = conv.blocks()
    doc.raw_text_chars = raw_text_chars(path, opts)
    return doc
