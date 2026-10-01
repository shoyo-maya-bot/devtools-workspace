"""PowerPoint (.pptx) → 中間表現（python-pptx）.

- スライドごとに H2（「N. タイトル」）
- 図形は上→下、左→右の順で読む（グループは展開）
- 本文プレースホルダの段落は箇条書き（段落レベルを保持）、その他のテキストボックスは段落
- 表・画像を抽出。発表者ノートは「ノート:」段落として末尾に付ける
- SmartArt の文字は `> **SmartArt**: …` として出力（図の形は再現しない）
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import defusedxml.ElementTree as SafeET
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..assets import AssetWriter
from ..model import Document, Heading, Image, ListItem, Options, Paragraph, Quote, Table, count_chars

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}


def _ordered_shapes(shapes):
    def key(s):
        return (s.top or 0, s.left or 0)

    for s in sorted(shapes, key=key):
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _ordered_shapes(s.shapes)
        else:
            yield s


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except ValueError:
        return None


A_T = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
DGM_DATA = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"


def raw_text_chars(path: Path) -> int:
    """zip 内のスライド・ノート・SmartArt の XML から a:t を直接数える（アダプタの解釈を通さない）."""
    total = 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if re.match(r"ppt/(slides/slide|notesSlides/notesSlide|diagrams/data)\d+\.xml$", name):
                root = SafeET.fromstring(z.read(name))
                if "notesSlide" in name:  # ノートのスライド番号プレースホルダ等は除く
                    texts = [t.text or "" for t in root.iter(A_T)]
                    total += count_chars("".join(texts)) - sum(count_chars(t) for t in texts if t.strip().isdigit())
                else:
                    total += count_chars("".join(t.text or "" for t in root.iter(A_T)))
    return total


def _smartart_texts(slide) -> list[str]:
    out = []
    for rel in slide.part.rels.values():
        if rel.reltype == DGM_DATA and not rel.is_external:
            root = SafeET.fromstring(rel.target_part.blob)
            texts = [t.text.strip() for t in root.iter(A_T) if t.text and t.text.strip()]
            if texts:
                out.append(" / ".join(texts))
    return out


def convert(path: Path, assets: AssetWriter, opts: Options | None = None) -> Document:
    prs = Presentation(str(path))
    doc = Document()
    st = doc.stats
    doc.blocks.append(Heading(1, path.stem, synthetic=True))
    for idx, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            st.add_text(title)
        st.headings += 1
        doc.blocks.append(Heading(2, f"{idx}. {title}" if title else f"{idx}."))
        for shape in _ordered_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            ph = _placeholder_type(shape)
            if shape.has_text_frame:
                as_list = ph in _BODY_TYPES
                for para in shape.text_frame.paragraphs:
                    text = para.text
                    if not text.strip():
                        continue
                    st.add_text(text)
                    if as_list:
                        st.list_items += 1
                        doc.blocks.append(ListItem(text, para.level, False))
                    else:
                        doc.blocks.append(Paragraph(text))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                st.tables += 1
                st.table_rows += sum(1 for r in rows if any(c.strip() for c in r))
                for r in rows:
                    for c in r:
                        st.add_text(c)
                doc.blocks.append(Table(rows=rows, header=True))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                st.images += 1
                rel = assets.write(img.blob, content_type=img.content_type, ext=img.ext)
                doc.blocks.append(Image(rel, alt=shape.name))
        for art in _smartart_texts(slide):
            st.add_text(art)
            doc.blocks.append(Quote(art, "SmartArt"))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                st.add_text(notes)
                doc.blocks.append(Paragraph("ノート: " + notes))
    doc.notes.append("グラフ・図形の形は対象外（SmartArt は文字のみ。必要なら --via-pdf を検討）")
    doc.raw_text_chars = raw_text_chars(path)
    return doc
