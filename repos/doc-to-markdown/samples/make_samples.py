"""サンプル文書を生成する（開発用。生成物はリポジトリにコミット済み）.

    pip install reportlab   # PDF 生成にのみ使用（ツール本体の依存ではない）
    python samples/make_samples.py

内容はすべて架空（経費精算の申請書という一般的な題材）。
"""

from __future__ import annotations

import base64
import datetime as dt
import urllib.parse
import zlib
from pathlib import Path

import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)

# 1x1 の PNG（抽出図のテスト用）
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
)


def make_docx() -> None:
    d = docx.Document()
    d.core_properties.created = d.core_properties.modified = FIXED_TIME
    d.add_heading("経費精算 申請書", level=0)
    d.add_heading("1. 概要", level=1)
    p = d.add_paragraph("本書は")
    p.add_run("出張旅費").bold = True
    p.add_run("の精算を申請するものです。")
    d.add_heading("1.1 対象", level=2)
    d.add_paragraph("交通費", style="List Bullet")
    d.add_paragraph("宿泊費", style="List Bullet")
    d.add_paragraph("日当", style="List Bullet 2")
    d.add_heading("2. 手順", level=1)
    for s in ["申請書を作成する", "上長が承認する", "経理が支払う"]:
        d.add_paragraph(s, style="List Number")
    d.add_heading("3. 明細", level=1)
    rows = [
        ("申請ID", "区分", "金額", "税額", "合計"),
        ("A-001", "交通費", "1000", "100", "1100"),
        ("A-002", "宿泊費", "8000", "800", "8800"),
    ]
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for r, row in enumerate(rows):
        for c, v in enumerate(row):
            t.cell(r, c).text = v
    d.add_paragraph("備考: 金額は税抜、合計は税込（単位: 円）。")
    img = HERE / "_tmp.png"
    img.write_bytes(PNG)
    d.add_picture(str(img))
    img.unlink()
    d.save(HERE / "expense-claim.docx")


def make_xlsx() -> None:
    wb = openpyxl.Workbook()
    wb.properties.created = wb.properties.modified = FIXED_TIME
    ws = wb.active
    ws.title = "申請一覧"
    ws.append(["申請ID", "申請日", "区分", "金額", "税額", "合計", "承認者", "部門コード"])
    data = [
        ("A-001", dt.date(2026, 4, 1), "交通費", 1000, 100, 1100, "佐藤", "D10"),
        ("A-002", dt.date(2026, 4, 2), "宿泊費", 8000, 800, 8900, "佐藤", "D10"),  # 合計が不一致
        ("A-003", dt.date(2026, 4, 3), "接待費", 30000, 3000, 33000, "", "D20"),  # 接待費なのに承認者なし
        ("A-004", dt.date(2026, 4, 5), "交通費", -500, 0, -500, "田中", "D99"),  # 負の金額・未登録部門
        ("A-005", dt.date(2026, 4, 8), "消耗品", 2400, 240, 2640, "田中", "D20"),
    ]
    for row in data:
        ws.append(row)
    ws2 = wb.create_sheet("部門マスタ")
    ws2.append(["部門コード", "部門名"])
    for row in [("D10", "営業部"), ("D20", "開発部"), ("D30", "管理部")]:
        ws2.append(row)
    wb.save(HERE / "expense-list.xlsx")


def make_pptx() -> None:
    prs = Presentation()
    prs.core_properties.created = prs.core_properties.modified = FIXED_TIME
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "経費精算の流れ"
    body = s.placeholders[1].text_frame
    body.text = "申請"
    for text, level in [("領収書を添付", 1), ("承認", 0), ("支払", 0)]:
        para = body.add_paragraph()
        para.text, para.level = text, level
    s.notes_slide.notes_text_frame.text = "承認は 3 営業日以内。"
    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "区分と上限"
    tbl = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.2)).table
    for r, row in enumerate([("区分", "上限"), ("交通費", "実費"), ("宿泊費", "10000")]):
        for c, v in enumerate(row):
            tbl.cell(r, c).text = v
    box = s2.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1)).text_frame
    box.text = "上限を超える場合は事前申請が必要です。"
    box.paragraphs[0].runs[0].font.size = Pt(14)
    img = HERE / "_tmp.png"
    img.write_bytes(PNG)
    s2.shapes.add_picture(str(img), Inches(7), Inches(4), Inches(0.5), Inches(0.5))
    img.unlink()
    prs.save(HERE / "expense-flow.pptx")


def make_pdf() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    c = canvas.Canvas(str(HERE / "expense-policy.pdf"), pagesize=A4, invariant=1)
    w, h = A4
    y = h - 72
    c.setFont("HeiseiKakuGo-W5", 20)
    c.drawString(72, y, "経費精算規程")
    y -= 40
    c.setFont("HeiseiKakuGo-W5", 10.5)
    for line in ["この規程は、業務上の経費の精算手続きを定める。", "申請は発生日から 30 日以内に行う。"]:
        c.drawString(72, y, line)
        y -= 16
    y -= 12
    c.setFont("HeiseiKakuGo-W5", 14)
    c.drawString(72, y, "第 2 条 対象")
    y -= 24
    c.setFont("HeiseiKakuGo-W5", 10.5)
    for line in ["・交通費", "・宿泊費", "・接待費（事前承認が必要）"]:
        c.drawString(72, y, line)
        y -= 16
    c.showPage()
    c.save()


GEOM = '<mxGeometry x="{x}" y="0" width="80" height="40" as="geometry"/>'
EDGE_GEOM = '<mxGeometry relative="1" as="geometry"/>'


def _drawio_xml(compressed: bool) -> str:
    model = (
        '<mxGraphModel><root><mxCell id="0"/><mxCell id="1" parent="0"/>'
        '<mxCell id="a" value="申請者" vertex="1" parent="1">' + GEOM.format(x=0) + "</mxCell>"
        '<mxCell id="b" value="上長&lt;br&gt;（承認）" vertex="1" parent="1">' + GEOM.format(x=160) + "</mxCell>"
        '<mxCell id="c" value="経理" vertex="1" parent="1">' + GEOM.format(x=320) + "</mxCell>"
        '<mxCell id="e1" value="申請" edge="1" parent="1" source="a" target="b">' + EDGE_GEOM + "</mxCell>"
        '<mxCell id="e2" edge="1" parent="1" source="b" target="c">' + EDGE_GEOM + "</mxCell>"
        "</root></mxGraphModel>"
    )
    if compressed:
        raw = urllib.parse.quote(model, safe="")
        co = zlib.compressobj(9, zlib.DEFLATED, -15)
        body = base64.b64encode(co.compress(raw.encode()) + co.flush()).decode()
        return f'<mxfile host="drawio"><diagram id="p1" name="承認フロー">{body}</diagram></mxfile>\n'
    return f'<mxfile host="drawio"><diagram id="p1" name="承認フロー">{model}</diagram></mxfile>\n'


def make_drawio() -> None:
    (HERE / "approval-flow.drawio").write_text(_drawio_xml(False), encoding="utf-8")
    (HERE / "approval-flow-compressed.drawio").write_text(_drawio_xml(True), encoding="utf-8")


if __name__ == "__main__":
    make_docx()
    make_xlsx()
    make_pptx()
    make_pdf()
    make_drawio()
    print("samples generated in", HERE)
