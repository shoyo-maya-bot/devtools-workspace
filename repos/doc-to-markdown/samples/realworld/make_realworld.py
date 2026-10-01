"""実務に近いサンプル文書を生成する（開発用。生成物はコミット済み）.

プログラムで直接作った文書は実務の文書より「きれい」になりがちなので、ここでは
**LibreOffice（実際のオフィスソフト）に保存させた** Word / Excel を作る。

    python samples/realworld/make_realworld.py        # 要 LibreOffice（DEVTOOLS_SOFFICE でも指定可）
    pip install reportlab pillow                        # PDF 系サンプルに使用

生成するもの（内容はすべて架空）:
- policy-revision.docx   変更履歴・コメント・脚注・テキストボックス・ヘッダー/フッター・縦結合・入れ子の表・ハイパーリンク
- expense-ledger.xlsx    タイトル行・複数行見出し（結合）・縦結合・1 シートに表 2 つ・非表示行・数式（計算済み）・注記
- uncalculated.xlsx      数式の計算結果が保存されていないブック（プログラム生成のブックでよくある）
- two-column.pdf         2 段組み・各ページに繰り返しのヘッダー/フッター
- scanned.pdf            画像だけの PDF（スキャン相当。英語。OCR の動作確認用）
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "devtools-common" / "src"))

NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:fo="urn:oasis:names:tc:opendocument:xmlns:xsl-fo-compatible:1.0" '
    'xmlns:xlink="http://www.w3.org/1999/xlink" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'xmlns:loext="urn:org:documentfoundation:names:experimental:office:xmlns:loext:1.0"'
)

DATE = "2026-04-01T10:00:00"


def change_info(who: str) -> str:
    return f"<office:change-info><dc:creator>{who}</dc:creator><dc:date>{DATE}</dc:date></office:change-info>"


def cell(text: str, extra: str = "") -> str:
    return f'<table:table-cell office:value-type="string" {extra}><text:p>{text}</text:p></table:table-cell>'


POLICY_FODT = f"""<?xml version="1.0" encoding="UTF-8"?>
<office:document {NS} office:version="1.3" office:mimetype="application/vnd.oasis.opendocument.text">
 <office:styles>
  <style:style style:name="Standard" style:family="paragraph"/>
  <style:style style:name="Heading" style:family="paragraph"/>
  <style:style style:name="Heading_20_1" style:display-name="Heading 1" style:family="paragraph"
     style:parent-style-name="Heading" style:default-outline-level="1">
   <style:text-properties fo:font-size="16pt" fo:font-weight="bold"/></style:style>
  <style:style style:name="Heading_20_2" style:display-name="Heading 2" style:family="paragraph"
     style:parent-style-name="Heading" style:default-outline-level="2">
   <style:text-properties fo:font-size="13pt" fo:font-weight="bold"/></style:style>
  <style:style style:name="Title" style:family="paragraph"><style:text-properties fo:font-size="20pt"/></style:style>
 </office:styles>
 <office:automatic-styles>
  <style:page-layout style:name="pm1"><style:page-layout-properties fo:page-width="21cm" fo:page-height="29.7cm"/>
   <style:header-style/><style:footer-style/></style:page-layout>
  <style:style style:name="T1" style:family="text"><style:text-properties fo:font-weight="bold"/></style:style>
  <text:list-style style:name="LB">
   <text:list-level-style-bullet text:level="1" text:bullet-char="•"/>
   <text:list-level-style-bullet text:level="2" text:bullet-char="◦"/>
  </text:list-style>
  <text:list-style style:name="LN">
   <text:list-level-style-number text:level="1" style:num-format="1" style:num-suffix="."/>
  </text:list-style>
 </office:automatic-styles>
 <office:master-styles>
  <style:master-page style:name="Standard" style:page-layout-name="pm1">
   <style:header><text:p>社外秘 経費精算規程 第2版</text:p></style:header>
   <style:footer><text:p>サンプル株式会社 総務部</text:p></style:footer>
  </style:master-page>
 </office:master-styles>
 <office:body><office:text>
  <text:tracked-changes>
   <text:changed-region text:id="ct1"><text:insertion>{change_info("佐藤")}</text:insertion></text:changed-region>
   <text:changed-region text:id="ct2"><text:deletion>{change_info("佐藤")}<text:p>8000 円</text:p></text:deletion>
   </text:changed-region>
  </text:tracked-changes>
  <text:p text:style-name="Title">経費精算規程</text:p>
  <text:h text:style-name="Heading_20_1" text:outline-level="1">第1条 目的</text:h>
  <text:p text:style-name="Standard">本規程は、業務上の経費の精算手続きを定める<text:note text:id="ftn1" text:note-class="footnote"><text:note-citation>1</text:note-citation><text:note-body><text:p>出張旅費規程を含む。</text:p></text:note-body></text:note>。詳細は<text:a xlink:type="simple" xlink:href="https://example.com/rules/expense">規程集</text:a>を参照する。</text:p>
  <text:h text:style-name="Heading_20_1" text:outline-level="1">第2条 対象と上限</text:h>
  <text:p text:style-name="Standard">宿泊費の上限は <text:change text:change-id="ct2"/><text:change-start text:change-id="ct1"/>10000 円<text:change-end text:change-id="ct1"/>とする。</text:p>
  <text:p text:style-name="Standard"><office:annotation office:name="c1"><dc:creator>田中</dc:creator><dc:date>{DATE}</dc:date><text:p>税込か税抜か明記してください</text:p></office:annotation>交通費は実費<office:annotation-end office:name="c1"/>を支給する。</text:p>
  <text:p text:style-name="Standard"><draw:frame draw:name="注意枠" text:anchor-type="paragraph" svg:width="12cm" svg:height="1.5cm" svg:x="0cm" svg:y="0cm"><draw:text-box><text:p>注意: 領収書の原本を提出すること</text:p></draw:text-box></draw:frame>区分ごとの扱いは次の表のとおり。</text:p>
  <table:table table:name="区分表">
   <table:table-column table:number-columns-repeated="3"/>
   <table:table-row>{cell("区分")}{cell("手段")}{cell("扱い")}</table:table-row>
   <table:table-row>{cell("交通費", 'table:number-rows-spanned="2"')}{cell("電車")}{cell("実費")}</table:table-row>
   <table:table-row><table:covered-table-cell/>{cell("タクシー")}{cell("事前承認")}</table:table-row>
   <table:table-row>{cell("宿泊費")}{cell("10000 円まで", 'table:number-columns-spanned="2"')}<table:covered-table-cell/></table:table-row>
   <table:table-row>{cell("接待費")}<table:table-cell office:value-type="string"><table:table table:name="内訳"><table:table-column table:number-columns-repeated="2"/><table:table-row>{cell("社内")}{cell("不可")}</table:table-row><table:table-row>{cell("社外")}{cell("可")}</table:table-row></table:table></table:table-cell>{cell("部長承認")}</table:table-row>
  </table:table>
  <text:h text:style-name="Heading_20_2" text:outline-level="2">申請の手順</text:h>
  <text:list text:style-name="LN">
   <text:list-item><text:p>申請書を作成する</text:p></text:list-item>
   <text:list-item><text:p>上長が<text:span text:style-name="T1">承認</text:span>する</text:p></text:list-item>
  </text:list>
  <text:list text:style-name="LB">
   <text:list-item><text:p>添付するもの</text:p>
    <text:list><text:list-item><text:p>領収書</text:p></text:list-item><text:list-item><text:p>行程表</text:p></text:list-item></text:list>
   </text:list-item>
  </text:list>
 </office:text></office:body>
</office:document>
"""


def soffice() -> str:
    from devtools_common.executables import SOFFICE, find_executable, not_found_message

    exe = find_executable(SOFFICE)
    if not exe:
        raise SystemExit(not_found_message(SOFFICE))
    return exe


def lo_convert(src: Path, fmt: str, out_dir: Path) -> Path:
    with tempfile.TemporaryDirectory() as prof:
        subprocess.run(
            [
                soffice(),
                f"-env:UserInstallation={Path(prof).as_uri()}",
                "--headless",
                "--convert-to",
                fmt,
                "--outdir",
                str(out_dir),
                str(src),
            ],
            check=True,
            capture_output=True,
            timeout=300,
        )
    return out_dir / f"{src.stem}.{fmt.split(':')[0]}"


def make_policy_docx() -> None:
    with tempfile.TemporaryDirectory() as d:
        src = Path(d) / "policy-revision.fodt"
        src.write_text(POLICY_FODT, encoding="utf-8")
        out = lo_convert(src, "docx:MS Word 2007 XML", Path(d))
        (HERE / "policy-revision.docx").write_bytes(out.read_bytes())


def make_ledger_xlsx() -> None:
    """openpyxl で作り、LibreOffice で開いて保存し直す（数式の計算結果がキャッシュされる = Excel で保存したのと同じ状態）."""
    import datetime as dt

    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4月"
    ws["A1"] = "経費精算台帳（2026年4月）"
    ws["A1"].font = Font(size=14, bold=True)
    ws.merge_cells("A1:G1")
    ws["A2"] = "作成: 総務部"
    # 2 段の見出し（「金額」が 3 列にまたがる）
    ws["A4"], ws["B4"], ws["C4"], ws["D4"], ws["G4"] = "申請ID", "申請日", "区分", "金額", "承認者"
    ws.merge_cells("A4:A5")
    ws.merge_cells("B4:B5")
    ws.merge_cells("C4:C5")
    ws.merge_cells("D4:F4")
    ws.merge_cells("G4:G5")
    ws["D5"], ws["E5"], ws["F5"] = "税抜", "税額", "税込"
    rows = [
        ("A-001", dt.date(2026, 4, 1), "交通費", 1000, None, "佐藤"),
        ("A-002", dt.date(2026, 4, 2), "宿泊費", 8000, None, "佐藤"),
        ("A-003", dt.date(2026, 4, 3), "接待費", 30000, None, None),
        ("A-004", dt.date(2026, 4, 5), "交通費", 2500, None, "田中"),
    ]
    for i, (aid, day, kind, net, _, who) in enumerate(rows, start=6):
        ws.cell(i, 1, aid)
        ws.cell(i, 2, day).number_format = "yyyy/mm/dd"
        ws.cell(i, 3, kind)
        ws.cell(i, 4, net)
        ws.cell(i, 5, f"=ROUNDDOWN(D{i}*0.1,0)")
        ws.cell(i, 6, f"=D{i}+E{i}")
        ws.cell(i, 7, who)
    ws.row_dimensions[8].hidden = True  # A-003 を非表示（実務でよくある「消したつもり」の行）
    ws["A10"] = "合計"
    ws["D10"], ws["E10"], ws["F10"] = "=SUM(D6:D9)", "=SUM(E6:E9)", "=SUM(F6:F9)"
    ws["A11"] = "※ 税額は 10% 端数切り捨て"
    # 同じシートの右側に 2 つ目の表（区分マスタ）
    ws["J4"], ws["K4"] = "区分", "上限"
    for r, (k, lim) in enumerate([("交通費", "実費"), ("宿泊費", 10000), ("接待費", "要承認")], start=5):
        ws.cell(r, 10, k)
        ws.cell(r, 11, lim)
    # 縦結合（部門ごとにまとめた表）
    ws2 = wb.create_sheet("部門別")
    ws2.append(["部門", "担当", "件数"])
    ws2.append(["営業部", "佐藤", 2])
    ws2.append([None, "鈴木", 1])
    ws2.append(["開発部", "田中", 1])
    ws2.merge_cells("A2:A3")
    with tempfile.TemporaryDirectory() as d:
        raw = Path(d) / "expense-ledger.xlsx"
        wb.save(raw)
        out_dir = Path(d) / "out"
        out_dir.mkdir()
        out = lo_convert(raw, "xlsx:Calc MS Excel 2007 XML", out_dir)
        (HERE / "expense-ledger.xlsx").write_bytes(out.read_bytes())
        # 比較用: LibreOffice を通さない = 数式の計算結果が無いブック
        wb2 = openpyxl.load_workbook(raw)
        wb2.save(HERE / "uncalculated.xlsx")


def make_pdfs() -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(UnicodeCIDFont("HeiseiKakuGo-W5"))
    font = "HeiseiKakuGo-W5"
    w, h = A4

    # 2 段組み + 繰り返しヘッダー/フッター
    c = canvas.Canvas(str(HERE / "two-column.pdf"), pagesize=A4, invariant=1)
    left_paras = [
        ("1. はじめに", ["本書は経費精算システムの運用手順を", "説明する。対象は全社員とする。"]),
        ("2. 申請", ["申請は発生日から 30 日以内に行う。", "領収書は原本を添付する。"]),
    ]
    right_paras = [
        ("3. 承認", ["承認者は 3 営業日以内に確認する。", "差し戻しの場合は理由を記載する。"]),
        ("4. 支払", ["承認後、翌月末に口座へ振り込む。", "立替額が 5 万円を超える場合は仮払を使う。"]),
    ]
    for page in (1, 2):
        c.setFont(font, 8)
        c.drawString(40, h - 30, "経費精算 運用手順書 第3版")
        c.drawRightString(w - 40, 30, f"{page} / 2")
        c.setFont(font, 18)
        c.drawString(40, h - 70, "経費精算 運用手順" if page == 1 else "経費精算 運用手順（続き）")
        for x, paras in ((40, left_paras), (w / 2 + 10, right_paras)):
            y = h - 110
            for title, lines in paras:
                c.setFont(font, 13)
                c.drawString(x, y, title if page == 1 else title.replace(".", "-" + str(page) + "."))
                y -= 20
                c.setFont(font, 10)
                for ln in lines:
                    c.drawString(x, y, ln)
                    y -= 14
                y -= 12
        c.showPage()
    c.save()

    # スキャン相当: 文字を画像として貼っただけの PDF（テキスト層なし）
    from PIL import Image, ImageDraw, ImageFont

    img = Image.new("L", (1240, 1754), 255)
    draw = ImageDraw.Draw(img)
    try:
        big = ImageFont.truetype("DejaVuSans-Bold.ttf", 44)
        body = ImageFont.truetype("DejaVuSans.ttf", 30)
    except OSError:
        big = body = ImageFont.load_default()
    draw.text((100, 120), "Expense Report", font=big, fill=0)
    for i, line in enumerate(
        ["Employee: Taro Yamada", "Date: 2026-04-01", "Total amount: 12,500 JPY", "Approved by: Hanako Sato"]
    ):
        draw.text((100, 240 + i * 60), line, font=body, fill=0)
    with tempfile.TemporaryDirectory() as d:
        png = Path(d) / "scan.png"
        img.save(png, dpi=(150, 150))
        c = canvas.Canvas(str(HERE / "scanned.pdf"), pagesize=A4, invariant=1)
        c.drawImage(str(png), 0, 0, width=w, height=h)
        c.showPage()
        c.save()


if __name__ == "__main__":
    make_policy_docx()
    make_ledger_xlsx()
    make_pdfs()
    print("generated:", ", ".join(sorted(p.name for p in HERE.iterdir() if p.suffix != ".py")))
