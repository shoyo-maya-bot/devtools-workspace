"""実務に近い文書（samples/realworld/ = LibreOffice に保存させた Word/Excel、2 段組み・スキャン PDF）のテスト.

ゴールデン: tests/golden/realworld/*.md（`UPDATE_GOLDEN=1 pytest` で更新し、差分を PR でレビュー）
外部ツールが要るテスト（LibreOffice・Tesseract）は、無い環境では skip する。
"""

from __future__ import annotations

import copy
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import docx
from devtools_common.executables import TESSERACT, find_executable
from devtools_common.report import Severity
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from doc_to_markdown.adapters import docx_adapter, drawio_adapter
from doc_to_markdown.converter import convert_file
from doc_to_markdown.model import Options
from doc_to_markdown.office import find_soffice

ROOT = Path(__file__).resolve().parent.parent
REAL = ROOT / "samples" / "realworld"
GOLDEN = Path(__file__).resolve().parent / "golden" / "realworld"

# (入力, オプション, ゴールデン名) — OCR は Tesseract の版で結果が揺れるのでゴールデンにしない
GOLDEN_CASES = [
    ("policy-revision.docx", {}, "policy-revision"),
    ("policy-revision.docx", {"track_changes": "show"}, "policy-revision.show"),
    ("expense-ledger.xlsx", {}, "expense-ledger"),
    ("expense-ledger.xlsx", {"skip_hidden": True}, "expense-ledger.skip-hidden"),
    ("uncalculated.xlsx", {}, "uncalculated"),
    ("two-column.pdf", {}, "two-column"),
]


def has_tesseract_lang(lang: str) -> bool:
    exe = find_executable(TESSERACT)
    if not exe:
        return False
    r = subprocess.run([exe, "--list-langs"], capture_output=True, text=True)
    return lang in r.stdout.split()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        p = mock.patch.object(drawio_adapter, "find_drawio_cli", return_value=None)
        p.start()
        self.addCleanup(p.stop)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def convert(self, name: str, **opts):
        r = convert_file(REAL / name, self.tmp / name.replace(".", "_"), options=Options(**opts))
        return r, r.output.read_text(encoding="utf-8")

    @staticmethod
    def rules(r, severity=None) -> list[str]:
        return [f.rule_id for f in r.findings if severity is None or f.severity is severity]


class GoldenTest(Base):
    def test_golden(self):
        for name, opts, golden in GOLDEN_CASES:
            with self.subTest(golden):
                r, md = self.convert(name, **opts)
                path = GOLDEN / f"{golden}.md"
                if os.environ.get("UPDATE_GOLDEN"):
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_text(md, encoding="utf-8", newline="\n")
                self.assertEqual(md, path.read_text(encoding="utf-8"))
                self.assertNotIn("loss.raw_text", self.rules(r), "元ファイルの文字が出力から欠けている")


class WordTest(Base):
    def test_everything_that_python_docx_alone_would_drop(self):
        r, md = self.convert("policy-revision.docx")
        self.assertIn("[規程集](https://example.com/rules/expense)", md)  # ハイパーリンク
        self.assertIn("宿泊費の上限は 10000 円とする。", md)  # 変更履歴（承認後）
        self.assertIn("> **テキストボックス**: 注意: 領収書の原本を提出すること", md)
        self.assertIn("| 交通費 | タクシー | 事前承認 |", md)  # 縦結合は値を繰り返す
        self.assertIn("社内 / 不可<br>社外 / 可", md)  # 入れ子の表
        self.assertIn("[^1]: 出張旅費規程を含む。", md)  # 脚注
        self.assertIn("[^c1]: コメント（田中）: 税込か税抜か明記してください（対象:「交通費は実費」）", md)
        self.assertIn("- ヘッダー: 社外秘 経費精算規程 第2版", md)
        self.assertIn("convert.warning", self.rules(r))  # 未確定の変更履歴は警告

    def test_track_changes_modes(self):
        _, reject = self.convert("policy-revision.docx", track_changes="reject")
        self.assertIn("宿泊費の上限は 8000 円とする。", reject)
        _, show = self.convert("policy-revision.docx", track_changes="show")
        self.assertIn("~~8000 円~~<ins>10000 円</ins>", show)

    def test_options_to_omit(self):
        _, md = self.convert("policy-revision.docx", include_comments=False, include_header_footer=False)
        self.assertNotIn("[^c", md)
        self.assertNotIn("ヘッダー", md)

    def test_content_control_is_read(self):
        """コンテンツコントロール（w:sdt）で囲まれた段落も落とさない."""
        d = docx.Document()
        d.add_paragraph("前")
        p = d.add_paragraph("コントロール内の本文")
        sdt = parse_xml(f'<w:sdt xmlns:w="{docx_adapter.W_NS}"><w:sdtPr/><w:sdtContent/></w:sdt>')
        p._p.addprevious(sdt)
        sdt.find(qn("w:sdtContent")).append(copy.deepcopy(p._p))
        p._p.getparent().remove(p._p)
        src = self.tmp / "sdt.docx"
        d.save(src)
        md = convert_file(src, self.tmp / "o").output.read_text(encoding="utf-8")
        self.assertIn("コントロール内の本文", md)

    def test_raw_check_detects_unsupported_element(self):
        """アダプタが表を読み飛ばした場合（未対応要素の模擬）、独立カウントで検知できること."""
        with mock.patch.object(docx_adapter._Converter, "table", lambda self, tbl: docx_adapter.Paragraph("")):
            r, _ = self.convert("policy-revision.docx")
        self.assertIn("loss.raw_text", self.rules(r))


class ExcelTest(Base):
    def test_real_world_layout(self):
        r, md = self.convert("expense-ledger.xlsx")
        self.assertIn("経費精算台帳（2026年4月）\n\n作成: 総務部", md)  # タイトル行は段落
        self.assertIn("| 申請ID | 申請日 | 区分 | 金額/税抜 | 金額/税額 | 金額/税込 | 承認者 |", md)  # 2 段見出し
        self.assertIn("| A-002 | 2026-04-02 | 宿泊費 | 8000 | 800 | 8800 | 佐藤 |", md)  # 数式の計算結果
        self.assertIn("### 表2（J4:K7）", md)  # 同じシートの 2 つ目の表
        self.assertIn("※ 税額は 10% 端数切り捨て", md)  # 注記は表の外
        self.assertIn("| 営業部 | 鈴木 | 1 |", md)  # 縦結合を埋める
        self.assertIn("convert.warning", self.rules(r))  # 非表示行の警告

    def test_skip_hidden(self):
        _, md = self.convert("expense-ledger.xlsx", skip_hidden=True)
        self.assertNotIn("A-003", md)
        self.assertIn("| A-004 |", md)

    def test_uncached_formulas_warned(self):
        r, md = self.convert("uncalculated.xlsx")
        msgs = [f.message for f in r.findings if f.severity is Severity.WARNING]
        self.assertTrue(any("計算結果が保存されていない" in m for m in msgs))
        self.assertIn("| A-001 | 2026-04-01 | 交通費 | 1000 |  |  | 佐藤 |", md)

    @unittest.skipUnless(find_soffice(), "LibreOffice not installed")
    def test_recalc(self):
        r, md = self.convert("uncalculated.xlsx", recalc=True)
        self.assertIn("| A-001 | 2026-04-01 | 交通費 | 1000 | 100 | 1100 | 佐藤 |", md)
        self.assertTrue(md.startswith("# uncalculated\n"))

    def test_header_row_override(self):
        _, md = self.convert("expense-ledger.xlsx", header_rows={"部門別": 2})
        self.assertIn("| 営業部 | 佐藤 | 2 |\n| --- | --- | --- |", md)


class PdfTest(Base):
    def test_two_columns_and_running_headers(self):
        r, md = self.convert("two-column.pdf")
        order = [md.index(s) for s in ["### 1. はじめに", "### 2. 申請", "### 3. 承認", "### 4. 支払"]]
        self.assertEqual(order, sorted(order), "左段 → 右段の順に読む")
        self.assertIn("本書は経費精算システムの運用手順を説明する。", md)  # 折り返しを結合
        self.assertNotIn("運用手順書 第3版", md)  # 繰り返しヘッダー
        self.assertNotIn("1 / 2", md)  # ページ番号

    def test_keep_running_headers(self):
        _, md = self.convert("two-column.pdf", remove_running_headers=False)
        self.assertIn("運用手順書 第3版", md)

    @unittest.skipUnless(has_tesseract_lang("eng"), "Tesseract (eng) not installed")
    def test_ocr_scanned_page(self):
        r, md = self.convert("scanned.pdf", ocr_lang="eng")
        self.assertIn("Expense Report", md)
        self.assertIn("12,500", md)
        self.assertIn("convert.warning", self.rules(r))  # OCR したことを必ず知らせる

    def test_ocr_never_warns_about_empty_page(self):
        r, md = self.convert("scanned.pdf", ocr="never")
        msgs = " ".join(f.message for f in r.findings)
        self.assertIn("文字の無いページ", msgs)

    def test_missing_ocr_language_is_reported(self):
        if not find_executable(TESSERACT):
            self.skipTest("Tesseract not installed")
        r, _ = self.convert("scanned.pdf", ocr_lang="zzz+eng")
        self.assertTrue(any("zzz" in f.message for f in r.findings))


if __name__ == "__main__":
    unittest.main()
